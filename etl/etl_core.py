# -*- coding: utf-8 -*-
"""
etl_core.py — общий модуль для ETL-пайплайнов контакт-центра.

Лежит в папке etl/ вместе с ноутбуками и скриптами, которые его
импортируют: monitoring.ipynb, alpha_registry.ipynb, omega_registry.ipynb
и остальными.

Здесь только то, что одинаково для всех проектов:
  - пути к данным и настройки из .env,
  - чтение Excel и определение даты отчёта,
  - подключение к Google Sheets,
  - поиск/создание листа по дате,
  - сериализация значений в JSON-совместимые типы,
  - форматирование листа через Sheets API,
  - демо-режим: Excel-книга вместо таблицы Google (DEMO=1 в .env).

Специфика каждого проекта живёт в своём ноутбуке.
"""

import logging
import os
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

import gspread
import numpy as np
import openpyxl
import pandas as pd
from dotenv import load_dotenv
from openpyxl.utils import range_boundaries

# ==========================================================================
# ПУТИ И НАСТРОЙКИ ОКРУЖЕНИЯ
# ==========================================================================

# Корень репозитория: etl_core.py лежит в etl/, значит на уровень выше.
ROOT = Path(__file__).resolve().parent.parent

# Выгрузки CRM и АТС. Папка в .gitignore: реальные данные в git не попадают.
DATA_DIR = ROOT / "data"

# Ключ сервисного аккаунта Google. Тоже в .gitignore.
CREDENTIALS = str(ROOT / "credentials.json")

# ID таблиц, адрес CRM и логины живут в .env (образец — .env.example).
load_dotenv(ROOT / ".env")

# DEMO=1 в .env — писать в Excel-файлы вместо Google Sheets и не ходить
# в CRM. Как это устроено — раздел «ДЕМО-РЕЖИМ» в конце файла.
DEMO = os.getenv("DEMO") == "1"
DEMO_DIR = DATA_DIR / "demo_output"


def env(name: str, default: str = "") -> str:
    """Настройка из .env. В коде таких значений нет — только имена."""
    value = os.getenv(name, default)
    # В демо-режиме таблиц Google нет, и ID им не нужен. Пустой ID
    # заменяется именем Excel-файла, куда пойдёт запись:
    # MONITORING_SHEET_ID -> data/demo_output/monitoring.xlsx
    if DEMO and not value and name.endswith("_SHEET_ID"):
        return name[:-len("_SHEET_ID")].lower()
    return value


# ==========================================================================
# ЛОГИРОВАНИЕ
# ==========================================================================

def setup_logger(name: str = "etl") -> logging.Logger:
    """
    Настраивает логгер для вывода в консоль DataSpell.

    Повторный вызов не плодит дубли обработчиков — важно, потому что
    в ноутбуке ячейку с импортом часто запускают несколько раз подряд.
    """
    lg = logging.getLogger(name)
    if not lg.handlers:
        h = logging.StreamHandler()
        h.setFormatter(logging.Formatter(
            "%(asctime)s | %(levelname)-7s | %(message)s", datefmt="%H:%M:%S"
        ))
        lg.addHandler(h)
    lg.setLevel(logging.INFO)
    lg.propagate = False
    return lg


log = setup_logger()

# Названия месяцев для месячных таблиц («Alpha Август 2026»)
RU_MONTHS: Dict[int, str] = {
    1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель",
    5: "Май", 6: "Июнь", 7: "Июль", 8: "Август",
    9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь",
}


def month_title(template: str, d: date) -> str:
    """
    Подставляет месяц и год в шаблон названия таблицы.

    Нужно потому, что таблицы Alpha и Omega заводятся помесячно:
    1 августа появляется новая таблица с новым ID, и хардкод ID
    перестал бы работать. Название же предсказуемо.

    >>> month_title("Alpha {month} {year}", date(2026, 8, 5))
    'Alpha Август 2026'
    """
    return template.format(month=RU_MONTHS[d.month], year=d.year)


# Часовая сетка отчётов: 00:00-01:00 ... 23:00-00:00
HOURS: List[int] = list(range(24))
TIME_LABELS: Dict[int, str] = {
    h: f"{h:02d}:00-{(h + 1) % 24:02d}:00" for h in HOURS
}


# ==========================================================================
# РАБОТА С ДАТАМИ
# ==========================================================================

class DateResolver:
    """
    Определяет дату отчёта и формирует имя листа Google Sheets.

    Дата берётся ИЗ ДАННЫХ, а не из системного времени: вы закрываете
    отчёты последовательно и можете делать 23-е число, когда на календаре
    уже 25-е. Единственный надёжный источник — колонка со временем звонка.
    """

    SHEET_FORMAT = "%d.%m.%y"          # основной формат имени листа: 23.07.26
    ALT_FORMATS = ["%d.%m.%Y"]         # старые листы вида 02.07.2026

    @staticmethod
    def from_dataframe(df: pd.DataFrame, date_column: str,
                       dayfirst: bool = False) -> date:
        """
        Извлекает дату отчёта из колонки с датой-временем.

        Если в файле несколько дат — берёт самую частую и предупреждает.
        Это защищает от ситуации, когда в выгрузку случайно попали
        соседние сутки: отчёт всё равно соберётся по основному дню.
        """
        if date_column not in df.columns:
            raise KeyError(f"В файле нет колонки '{date_column}'. "
                           f"Фактические: {list(df.columns)}")

        # dayfirst нужен для колонок вида «05.07.2026»: без него
        # пятое июля можно прочитать как седьмое мая
        dt = pd.to_datetime(df[date_column], errors="coerce",
                            dayfirst=dayfirst).dropna()
        if dt.empty:
            if df.empty:
                raise ValueError(
                    "Файл пустой — в выгрузке ни одной строки. "
                    "Либо за этот день звонков не было, либо выгрузка не удалась."
                )
            raise ValueError(
                f"Колонка '{date_column}' не содержит распознаваемых дат "
                f"(строк в файле: {len(df)})"
            )

        days = dt.dt.date.value_counts()
        target = days.index[0]

        if len(days) > 1:
            log.warning(
                "В файле несколько дат: %s. Беру основную: %s",
                dict(list(days.items())[:5]), target.strftime("%d.%m.%Y")
            )
        return target

    @classmethod
    def sheet_name(cls, d: date) -> str:
        """Имя листа в основном формате — 23.07.26."""
        return d.strftime(cls.SHEET_FORMAT)

    @classmethod
    def parse_sheet_name(cls, title: str) -> Optional[date]:
        """
        Обратная операция: имя листа -> дата. Не-дата возвращает None.

        Нужно для сбора истории: среди вкладок таблицы попадаются
        «Итог Июль» и прочие, их надо отсеять.
        """
        t = str(title).strip()
        for f in [cls.SHEET_FORMAT] + cls.ALT_FORMATS:
            try:
                return datetime.strptime(t, f).date()
            except ValueError:
                continue
        return None

    @classmethod
    def candidate_names(cls, d: date) -> List[str]:
        """
        Все варианты написания имени листа для этой даты.

        Нужно, потому что часть старых листов названа полным годом
        (02.07.2026). При поиске проверяем оба, чтобы не создать дубль.
        """
        return [d.strftime(f) for f in [cls.SHEET_FORMAT] + cls.ALT_FORMATS]


# ==========================================================================
# ЧТЕНИЕ ИСХОДНИКОВ
# ==========================================================================

class ExcelReader:
    """Чтение выгрузок CRM и АТС с проверкой структуры."""

    # Ожидаемые колонки CRM — одинаковы для всех проектов
    CRM_REQUIRED = ["ID звонка", "Направление звонка", "Время начала звонка",
                    "Продолжительность, сек", "Оператор"]

    def __init__(self, base_dir: Path):
        """:param base_dir: папка с Excel-файлами (обычно папка data/)"""
        self.base_dir = Path(base_dir)

    def _path(self, filename: str) -> Path:
        """Собирает полный путь и падает сразу, если файла нет."""
        p = self.base_dir / filename
        if not p.exists():
            raise FileNotFoundError(
                f"Файл не найден: {p}\n"
                f"Проверьте, что выгрузка лежит в {self.base_dir}"
            )
        return p

    def read_crm(self, filename: str) -> pd.DataFrame:
        """
        Читает CRM-выгрузку (.xlsx) с проверкой обязательных колонок.

        Колонки берутся ПО ИМЕНИ, а не по позиции: если CRM добавит
        поле слева, скрипт продолжит читать правильные данные,
        а не сдвинется молча на соседний столбец.
        """
        p = self._path(filename)
        df = pd.read_excel(p)
        df.columns = [str(c).strip() for c in df.columns]

        missing = [c for c in self.CRM_REQUIRED if c not in df.columns]
        if missing:
            raise KeyError(f"{filename}: нет обязательных колонок {missing}")

        log.info("CRM %s: %d строк", filename, len(df))
        return df

    def read_pbx_registry(self, filename: str) -> pd.DataFrame:
        """
        Читает выгрузку АТС в формате РЕЕСТРА ЗВОНКОВ (проект Alpha).

        Шапка занимает две строки, поэтому skiprows=2 и колонки по индексам.
        Возвращает нормализованный DataFrame с понятными именами.
        """
        p = self._path(filename)
        raw = pd.read_excel(p, header=None, skiprows=2)

        if raw.shape[1] < 8:
            raise ValueError(
                f"{filename}: ожидался реестр звонков (>=8 колонок), "
                f"получено {raw.shape[1]}. Похоже, это почасовой отчёт — "
                f"используйте read_pbx_hourly()"
            )

        df = pd.DataFrame({
            "Дата": raw[0],
            "Время": raw[1],
            "Результат": raw[2].astype(str).str.strip(),
            "Линия": raw[3],
            # Номер абонента приходит с ведущим «+» — убираем,
            # иначе Sheets покажет «++998...» после сериализации
            "Номер": raw[4].astype(str).str.lstrip("+").str.strip(),
            "IVR_длит": raw[5],
            "Внутренний": raw[6].astype(str).str.strip(),
            "Оператор": raw[7],
            "Длит_оператора": raw[8],
            # Колонка 20 — «Продолжительность / Итого», крайняя справа
            "Итого": raw[20] if raw.shape[1] > 20 else "",
        })
        log.info("АТС %s: %d звонков", filename, len(df))
        return df

    @staticmethod
    def _to_num(series: pd.Series) -> pd.Series:
        """
        Числа из выгрузки АТС, устойчиво к разделителю разрядов.

        Начиная с тысячи значения приходят как «1 036», иногда
        с неразрывным пробелом. Обычный to_numeric на такой строке даёт
        NaN, а fillna(0) молча превращает тысячу с лишним звонков в ноль.
        """
        s = (series.astype(str)
             .str.replace("\xa0", "", regex=False)
             .str.replace("\u2009", "", regex=False)
             .str.replace(" ", "", regex=False)
             .str.replace(",", ".", regex=False))
        return pd.to_numeric(s, errors="coerce")

    @staticmethod
    def _detect_pbx_format(raw: pd.DataFrame) -> str:
        """
        Определяет, что за файл прислала АТС.

        Возвращает "report" — готовый почасовой отчёт (24 строки + Итого),
        или "registry" — реестр звонков (строка на звонок).

        Различаются по первой колонке: в отчёте там время «0:00:00»,
        в реестре — дата «23.07.2026». Без этой проверки дата
        разбирается как час (23) для всех строк сразу.
        """
        col0 = raw[0].astype(str)
        if col0.str.match(r"^\s*\d{1,2}\.\d{1,2}\.\d{2,4}").any():
            return "registry"
        # Запасной признак: колонка «Результат» со значениями реестра
        if raw.shape[1] > 2 and raw[2].astype(str).str.contains(
                "Обработан|Потерян", na=False).any():
            return "registry"
        return "report"

    @staticmethod
    def _hour_from_time(series: pd.Series) -> pd.Series:
        """Достаёт час из ячейки времени. Нераспознанное -> NaN."""
        h = pd.to_datetime(series.astype(str).str.strip(),
                           format="%H:%M:%S", errors="coerce").dt.hour
        # Запасной разбор для формата без секунд или с датой впереди
        miss = h.isna()
        if miss.any():
            alt = series[miss].astype(str).str.extract(
                r"(?:^|\s)(\d{1,2}):\d{2}")[0]
            h.loc[miss] = pd.to_numeric(alt, errors="coerce")
        return h

    def _hourly_from_registry(self, raw: pd.DataFrame,
                              filename: str) -> pd.DataFrame:
        """
        Сворачивает реестр звонков в почасовые агрегаты.

        Соответствие показателей то же, что у Alpha:
            «Обработан IVR»                -> IVR
            «Обработан первым оператором»  -> Входящие
            «Потерян»                      -> Потерянные
        """
        res = raw[2].astype(str).str.strip()
        h = self._hour_from_time(raw[1])

        df = pd.DataFrame({"Час": h, "Результат": res})
        bad = int(df["Час"].isna().sum())
        if bad:
            log.warning("%s: %d строк с нераспознанным временем", filename, bad)
        df = df.dropna(subset=["Час"])
        df["Час"] = df["Час"].astype(int)

        def by(mask_value: str) -> pd.Series:
            return (df[df["Результат"] == mask_value]
                    .groupby("Час").size().reindex(HOURS, fill_value=0))

        out = pd.DataFrame({
            "Час": HOURS,
            "IVR": by("Обработан IVR").values,
            "Входящие": by("Обработан первым оператором").values,
            "Потерянные": by("Потерян").values,
        })

        known = int(out[["IVR", "Входящие", "Потерянные"]].values.sum())
        if known != len(df):
            other = sorted(set(df["Результат"]) - {
                "Обработан IVR", "Обработан первым оператором", "Потерян"})
            log.warning("%s: %d строк с прочими результатами не учтены: %s",
                        filename, len(df) - known, other)

        log.info("%s: реестр свёрнут в почасовые агрегаты "
                 "(IVR=%d, входящие=%d, потерянные=%d)",
                 filename, out["IVR"].sum(), out["Входящие"].sum(),
                 out["Потерянные"].sum())
        return out

    def _hourly_from_report(self, raw: pd.DataFrame,
                            filename: str) -> pd.DataFrame:
        """
        Читает готовый почасовой отчёт АТС.

        Колонки блока «Количество звонков»:
            0 Время | 1 IVR | 2 Первый оператор | 3 Операторы
            4 Внешний агент | 5 Потеряно | 6 Итого
        """
        if raw.shape[1] < 7:
            raise ValueError(f"{filename}: слишком мало колонок ({raw.shape[1]})")

        # Отсекаем строку «Итого» — она не относится к часовой сетке
        mask_total = raw[0].astype(str).str.contains("Итог", case=False, na=False)
        body = raw[~mask_total].copy()

        def to_hour(v: Any) -> Optional[int]:
            """'14:00:00' -> 14. Возвращает None для нераспознанного."""
            m = re.match(r"^\s*(\d{1,2})[:.]", str(v))
            return int(m.group(1)) if m else None

        body["Час"] = body[0].apply(to_hour)
        body = body.dropna(subset=["Час"])
        body["Час"] = body["Час"].astype(int)

        out = pd.DataFrame({
            "Час": body["Час"].values,
            "IVR": self._to_num(body[1]).fillna(0).astype(int).values,
            "Входящие": self._to_num(body[2]).fillna(0).astype(int).values,
            "Потерянные": self._to_num(body[5]).fillna(0).astype(int).values,
        })
        return out

    @staticmethod
    def _registry_diagnosis(raw: pd.DataFrame, filename: str) -> str:
        """Собирает описание содержимого реестра для сообщения об ошибке."""
        dates = sorted(set(raw[0].astype(str).str.strip()))[:3]
        lines = raw[3].value_counts().to_dict() if raw.shape[1] > 3 else {}
        res = raw[2].value_counts().to_dict() if raw.shape[1] > 2 else {}
        internal = raw[6].value_counts().to_dict() if raw.shape[1] > 6 else {}
        return (
            f"\n  Файл:         {filename}"
            f"\n  Формат:       реестр звонков ({len(raw)} строк, "
            f"{raw.shape[1]} колонок)"
            f"\n  Дата внутри:  {', '.join(dates)}"
            f"\n  Номер линии:  {lines}"
            f"\n  Результаты:   {res}"
            f"\n  Внутренние:   {internal}"
        )

    def read_pbx_hourly(self, filename: str,
                        allow_registry: bool = False) -> pd.DataFrame:
        """
        Возвращает почасовые агрегаты АТС: Час, IVR, Входящие, Потерянные.

        Для проектов «Мониторинга» АТС всегда отдаёт ГОТОВЫЙ ПОЧАСОВОЙ ОТЧЁТ.
        Если вместо него пришёл реестр звонков — это почти наверняка
        значит, что выгружен не тот отчёт или файл от другого проекта.
        Поэтому по умолчанию скрипт останавливается и показывает, что
        внутри файла: свернуть реестр молча опаснее, чем упасть —
        чужие цифры ушли бы в таблицу и остались бы там незамеченными.

        :param allow_registry: True — всё-таки свернуть реестр по часам.
            Включать, только если формат выгрузки сменился намеренно.
        """
        p = self._path(filename)
        raw = pd.read_excel(p, header=None, skiprows=2)

        fmt = self._detect_pbx_format(raw)
        if fmt == "registry":
            if not allow_registry:
                raise ValueError(
                    "АТС прислала РЕЕСТР звонков вместо почасового отчёта."
                    + self._registry_diagnosis(raw, filename) +
                    "\n\n  Что делать:"
                    "\n  1. Сверьте дату и номер линии — файл точно от этого проекта?"
                    "\n  2. Перевыгрузите отчёт в том же виде, что у остальных проектов"
                    "\n  3. Если формат сменился намеренно — поставьте"
                    "\n     ALLOW_REGISTRY = True в ячейке 1"
                )
            log.warning("%s: реестр вместо почасового отчёта, "
                        "сворачиваю по часам (ALLOW_REGISTRY = True)", filename)
            out = self._hourly_from_registry(raw, filename)
        else:
            out = self._hourly_from_report(raw, filename)

        # Час обязан быть уникальным: на дубликатах падает reindex,
        # и причина ошибки становится неочевидной
        if not out["Час"].is_unique:
            dup = out.loc[out["Час"].duplicated(keep=False), "Час"].unique()
            raise ValueError(
                f"{filename}: часы повторяются {sorted(dup)}. "
                f"Похоже, файл содержит несколько дней или формат сменился"
            )

        log.info("АТС %s: %d часовых строк", filename, len(out))
        return out

# ==========================================================================
# ВЫГРУЗКА ЗА НЕСКОЛЬКО ДНЕЙ
# ==========================================================================

class MultiDay:
    """
    Разбор выгрузки, охватывающей несколько дней.

    Зачем это нужно. Выгрузка CRM теряет ПЕРВУЮ запись периода — ровно
    одну, независимо от длины периода. При суточной выгрузке это первый
    звонок дня, и он пропадает каждый день. При выгрузке за несколько
    дней теряется только первая запись самого раннего дня.

    Отсюда приём: берём период с запасом в один день назад. Первый день
    считается буферным — он показывается на экране, но не загружается,
    потому что именно у него откушена первая запись. Остальные дни
    приходят целыми.

    Если в файле один день, буфера нет и первая запись потеряна —
    об этом говорится прямо.
    """

    @staticmethod
    def split(df: pd.DataFrame, date_column: str,
              dayfirst: bool = False) -> Dict[date, pd.DataFrame]:
        """
        Раскладывает таблицу по датам. Строки без даты отбрасываются.

        :param dayfirst: True для колонок вида «05.07.2026»
        """
        if date_column not in df.columns:
            raise KeyError(f"Нет колонки '{date_column}'. "
                           f"Есть: {list(df.columns)}")

        dt = pd.to_datetime(df[date_column], errors="coerce", dayfirst=dayfirst)
        bad = int(dt.isna().sum())
        if bad:
            log.warning("Строк с нераспознанной датой отброшено: %d", bad)

        out = {}
        for d, part in df.assign(_d=dt.dt.date).dropna(subset=["_d"]).groupby("_d"):
            out[d] = part.drop(columns=["_d"])
        return dict(sorted(out.items()))

    @staticmethod
    def plan(parts: Dict[date, pd.DataFrame],
             use_buffer: bool = True) -> tuple:
        """
        Решает, какие дни грузить, а какой считать буферным.

        Возвращает (таблица для показа, список дат к загрузке, буферный день).
        Ничего не отбрасывает молча: буферный день попадает в таблицу
        с пометкой, чтобы его было видно перед запуском выгрузки.
        """
        days = sorted(parts)
        if not days:
            raise ValueError("В файле не нашлось ни одной строки с датой")

        buffer_day = days[0] if (use_buffer and len(days) > 1) else None
        to_load = [d for d in days if d != buffer_day]

        rows = []
        for d in days:
            if d == buffer_day:
                verdict = "буферный — НЕ грузится"
                note = "у него откушена первая запись, он и нужен как жертва"
            elif len(days) == 1:
                verdict = "грузится"
                note = ("буфера нет: в файле один день, первая запись потеряна. "
                        "Возьмите период на день шире")
            else:
                verdict = "грузится"
                note = ""
            rows.append({
                "Дата": d.strftime("%d.%m.%Y"),
                "День недели": WEEKDAYS_SHORT[d.weekday()],
                "Строк": len(parts[d]),
                "Что делаем": verdict,
                "Примечание": note,
            })

        return pd.DataFrame(rows), to_load, buffer_day


WEEKDAYS_SHORT = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]


# ==========================================================================
# МИНУТЫ ДЛЯ ТАРИФИКАЦИИ
# ==========================================================================

def bill_minutes(sec, free_under_sec: Optional[int] = None):
    """
    Секунды разговора -> минуты для тарификации.

    Обычно округление ВВЕРХ: разговор в 61 секунду — уже две минуты.
    Если заказчик задал порог (у Alpha — free_under_sec в
    reconciliation.ipynb), звонки короче или равные ему стоят 0 минут,
    а не округляются до одной.

    Принимает число или Series: работает и на одном звонке, и на всей
    выгрузке. Вынесено сюда из reconciliation.ipynb, чтобы правило
    можно было проверить тестом.
    """
    if free_under_sec:
        return np.where(sec <= free_under_sec, 0, np.ceil(sec / 60))
    return np.ceil(sec / 60)


# ==========================================================================
# СЕРИАЛИЗАЦИЯ
# ==========================================================================

class Serializer:
    """
    Конвертация значений в базовые типы Python для gspread.

    numpy.int64, numpy.float64 и pandas.Timestamp не сериализуются в JSON —
    Google API вернёт ошибку. Здесь всё приводится к int / float / str.
    """

    @staticmethod
    def cast(v: Any) -> Any:
        """Одно значение -> JSON-совместимый тип. Округления не происходит."""
        if v is None:
            return ""
        if isinstance(v, float) and np.isnan(v):
            return ""
        if isinstance(v, (np.integer,)):
            return int(v)
        if isinstance(v, (np.floating,)):
            f = float(v)
            # int только если число и так целое — значение не меняется
            return int(f) if f.is_integer() else f
        if isinstance(v, (np.bool_, bool)):
            return bool(v)
        if isinstance(v, (pd.Timestamp, datetime)):
            return v.strftime("%Y-%m-%d %H:%M:%S")
        if isinstance(v, (int, float, str)):
            return v
        if pd.isna(v):
            return ""
        return str(v)

    @classmethod
    def frame(cls, df: pd.DataFrame, header: bool = True) -> List[List[Any]]:
        """
        DataFrame -> список списков.

        Индекс НЕ выгружается: itertuples(index=False) гарантирует,
        что фантомный столбец с номерами строк в таблицу не попадёт.
        """
        rows = [[cls.cast(v) for v in r] for r in df.itertuples(index=False, name=None)]
        if header:
            return [[str(c) for c in df.columns]] + rows
        return rows


# ==========================================================================
# GOOGLE SHEETS
# ==========================================================================

# Палитра — подобрана под текущий вид ваших таблиц
class Palette:
    HEADER_GREEN = {"red": 0.78, "green": 0.87, "blue": 0.75}   # шапка светло-зелёная
    HEADER_DARK = {"red": 0.11, "green": 0.31, "blue": 0.21}    # шапка тёмно-зелёная
    ROW_YELLOW = {"red": 1.0, "green": 0.95, "blue": 0.80}      # подсветка активных часов
    TOTAL_YELLOW = {"red": 1.0, "green": 0.93, "blue": 0.70}    # строка «Итого»
    WHITE = {"red": 1.0, "green": 1.0, "blue": 1.0}
    BLACK = {"red": 0.0, "green": 0.0, "blue": 0.0}
    BORDER = {"red": 0.4, "green": 0.4, "blue": 0.4}


class SheetsLoader:
    """
    Выгрузка в Google Sheets: подключение, поиск листа по дате,
    запись значений и форматирование.
    """

    def __init__(self, credentials_path: str,
                 spreadsheet_id: Optional[str] = None):
        """
        :param credentials_path: путь к credentials.json сервисного аккаунта
        :param spreadsheet_id: ID таблицы из URL (часть между /d/ и /edit).
            Можно не задавать, если таблица открывается по названию —
            см. connect(title=...)
        """
        self.credentials_path = Path(credentials_path)
        self.spreadsheet_id = spreadsheet_id
        self.gc = None
        self.sh = None
        self.ws = None
        self.sheet_created = False

    def _pick_title(self, titles: List[str]) -> str:
        """
        Выбирает из вариантов написания тот, что реально существует.

        Нужно, когда название таблицы завели с опечаткой: скрипт
        примет и его, и исправленное, так что переименование
        таблицы не потребует правки кода.

        Если проверить не удалось (Drive API закрыт) — возвращает
        первый вариант, дальше сработает откат по ID.
        """
        if len(titles) == 1:
            return titles[0]
        try:
            existing = {s.title.strip().lower(): s.title
                        for s in self.gc.openall()}
        except Exception as e:
            log.warning("Список таблиц недоступен (%s) — беру первый вариант",
                        type(e).__name__)
            return titles[0]

        for t in titles:
            if t.strip().lower() in existing:
                if t != titles[0]:
                    log.info("Название таблицы: «%s»", t)
                return t
        return titles[0]

    def _open_by_title(self, title: str, attempts: int = 3):
        """
        Открывает таблицу по названию, переживая разовый обрыв связи.

        SSL-разрывы на этом запросе случаются и сами по себе, поэтому
        перед тем как считать Drive API недоступным, пробуем ещё раз.
        """
        import time
        last = None
        for i in range(attempts):
            try:
                return self.gc.open(title)
            except gspread.exceptions.SpreadsheetNotFound:
                raise
            except Exception as e:
                last = e
                if i < attempts - 1:
                    log.warning("Попытка %d/%d не удалась (%s), повтор",
                                i + 1, attempts, type(e).__name__)
                    time.sleep(2 * (i + 1))
        raise last

    def connect(self, title: Optional[str] = None):
        """
        Авторизация по сервисному аккаунту и открытие таблицы.

        :param title: если задано — таблица ищется ПО НАЗВАНИЮ.
            Так открываются месячные таблицы: 1 августа появляется
            «Alpha Август 2026» с новым ID, но название
            предсказуемо и правка кода не нужна.
            Если не задано — открывается по ID.

        При поиске по названию скрипт не подставляет запасной ID:
        иначе августовские данные молча ушли бы в июльскую таблицу.
        Вместо этого показывает список доступных таблиц.
        """
        if DEMO:
            # Вместо таблицы Google — Excel-книга с тем же названием
            # в data/demo_output. Ключ Google в демо не нужен.
            if title and not isinstance(title, str):
                title = list(title)[0]
            name = title or self.spreadsheet_id
            if not name:
                raise ValueError("Не задан ни spreadsheet_id, ни title")
            self.sh = DemoBook(name)
            log.info("Демо-режим: вместо Google Sheets — %s", self.sh.path)
            return self

        if not self.credentials_path.exists():
            raise FileNotFoundError(f"credentials.json не найден: {self.credentials_path}")

        self.gc = gspread.service_account(filename=str(self.credentials_path))

        if title:
            # title может быть списком: у таблицы бывает несколько
            # написаний названия, и все они считаются подходящими
            titles = [title] if isinstance(title, str) else list(title)
            title = self._pick_title(titles)

            try:
                self.sh = self._open_by_title(title)
            except gspread.exceptions.SpreadsheetNotFound:
                try:
                    available = sorted(s.title for s in self.gc.openall())
                    hint = "\nДоступные сейчас таблицы:\n  " + "\n  ".join(available)
                except Exception:
                    hint = ""
                raise RuntimeError(
                    f"Таблица «{title}» не найдена.\n"
                    f"Проверьте: создана ли она и открыт ли доступ "
                    f"сервисному аккаунту." + hint
                ) from None
            except Exception as e:
                # Поиск по названию идёт через Drive API (www.googleapis.com),
                # а данные — через Sheets API (sheets.googleapis.com). Первый
                # бывает закрыт файрволом, второй при этом работает.
                # Тогда открываем по ID, но обязательно сверяем название:
                # иначе августовские данные молча ушли бы в июльскую таблицу.
                if not self.spreadsheet_id:
                    raise
                log.warning("Поиск по названию недоступен (%s: %s). "
                            "Пробую открыть по ID",
                            type(e).__name__, str(e)[:120])
                self.sh = self.gc.open_by_key(self.spreadsheet_id)
                if self.sh.title.strip().lower() != title.strip().lower():
                    raise RuntimeError(
                        f"Таблица по ID называется «{self.sh.title}», "
                        f"а нужна «{title}».\n"
                        f"Поиск по названию сейчас недоступен, поэтому "
                        f"подставить нужную таблицу автоматически не выйдет.\n"
                        f"Впишите ID новой таблицы в ячейку 1."
                    ) from None
                log.info("Открыто по ID, название совпало")
        else:
            if not self.spreadsheet_id:
                raise ValueError("Не задан ни spreadsheet_id, ни title")
            self.sh = self.gc.open_by_key(self.spreadsheet_id)

        log.info("Подключено к таблице: %s", self.sh.title)
        return self

    def open_sheet_for_date(self, d: date, rows: int = 1000, cols: int = 30,
                            clear_range: Optional[str] = None):
        """
        Находит лист по дате или создаёт новый.

        Проверяются оба формата имени (23.07.26 и 23.07.2026), чтобы
        не создать дубль рядом с уже существующим листом.

        :param clear_range: если задан (например "A1:G200") — очищается
            только этот диапазон. Нужно, когда на листе есть колонки
            с ручными формулами, которые нельзя стирать.
            По умолчанию очищается весь лист.
        """
        existing = {w.title: w for w in self.sh.worksheets()}

        for name in DateResolver.candidate_names(d):
            if name in existing:
                self.ws = existing[name]
                if clear_range:
                    self.ws.batch_clear([clear_range])
                    log.info("Лист '%s' найден, очищен диапазон %s",
                             name, clear_range)
                else:
                    self.ws.clear()
                    log.info("Лист '%s' найден и очищен", name)
                self.sheet_created = False
                return self.ws

        name = DateResolver.sheet_name(d)
        self.ws = self.sh.add_worksheet(title=name, rows=rows, cols=cols)
        log.info("Создан новый лист '%s'", name)
        # Флаг нужен вызывающему коду: на новом листе надо заполнить
        # блоки, которые на существующих листах уже настроены вручную
        self.sheet_created = True
        return self.ws

    def open_sheet_by_name(self, name: str, rows: int = 3000, cols: int = 30):
        """Открывает лист с фиксированным именем (например 'raw data')."""
        try:
            self.ws = self.sh.worksheet(name)
            self.ws.clear()
            log.info("Лист '%s' найден и очищен", name)
        except gspread.exceptions.WorksheetNotFound:
            self.ws = self.sh.add_worksheet(title=name, rows=rows, cols=cols)
            log.info("Создан новый лист '%s'", name)
        return self.ws

    def reset_sheet_format(self) -> None:
        """
        Снимает с листа всё оформление, накопленное прошлыми запусками.

        clear() стирает только значения. Числовые форматы, заливки и
        правила условного форматирования остаются жить на листе, и при
        следующем запуске с другим набором колонок ложатся не туда:
        так штучная колонка получает процентный формат и показывает
        279 как 27900%. Поэтому перед каждой записью лист раздевается
        до чистого состояния.
        """
        sid = self._sheet_id()
        try:
            meta = self.sh.fetch_sheet_metadata()
            sheet = next(x for x in meta["sheets"]
                         if x["properties"]["sheetId"] == sid)
            n_rules = len(sheet.get("conditionalFormats", []))
        except Exception as e:
            log.warning("Не удалось прочитать правила листа (%s)",
                        type(e).__name__)
            n_rules = 0

        reqs = [{"deleteConditionalFormatRule": {"sheetId": sid, "index": i}}
                for i in range(n_rules - 1, -1, -1)]
        reqs.append({
            "repeatCell": {
                "range": {"sheetId": sid},
                "cell": {},
                "fields": "userEnteredFormat",
            }
        })
        try:
            self.sh.batch_update({"requests": reqs})
            if n_rules:
                log.info("Снято старых правил оформления: %d", n_rules)
        except Exception as e:
            log.warning("Сброс оформления не удался (%s)", type(e).__name__)

        # Объединённые ячейки убираются отдельно: на листе без
        # объединений этот запрос возвращает ошибку
        try:
            self.sh.batch_update({"requests": [
                {"unmergeCells": {"range": {"sheetId": sid}}}]})
        except Exception:
            pass

    def prune_sheets(self, keep: List[str]) -> None:
        """
        Удаляет вкладки, которых нет в списке нужных.

        Без этого вкладки от прежних версий отчёта остаются рядом
        с новыми и показывают устаревшие цифры — читающий не понимает,
        каким верить.
        """
        keep_set = {k.strip().lower() for k in keep}
        extra = [w for w in self.sh.worksheets()
                 if w.title.strip().lower() not in keep_set]
        if not extra or len(extra) >= len(self.sh.worksheets()):
            return
        for w in extra:
            try:
                self.sh.del_worksheet(w)
                log.info("Удалена устаревшая вкладка «%s»", w.title)
            except Exception as e:
                log.warning("Вкладку «%s» удалить не вышло (%s)",
                            w.title, type(e).__name__)

    def reorder_sheets(self, order: List[str]) -> None:
        """Расставляет вкладки в заданном порядке — слева направо."""
        titles = {w.title.strip().lower(): w for w in self.sh.worksheets()}
        reqs = []
        for i, name in enumerate(order):
            w = titles.get(name.strip().lower())
            if w is None:
                continue
            reqs.append({
                "updateSheetProperties": {
                    "properties": {"sheetId": w._properties["sheetId"],
                                   "index": i},
                    "fields": "index",
                }
            })
        if reqs:
            try:
                self.sh.batch_update({"requests": reqs})
                log.info("Вкладки упорядочены")
            except Exception as e:
                log.warning("Порядок вкладок не применился (%s)",
                            type(e).__name__)

    def ensure_month_table(self, want: str, previous: Optional[str] = None):
        """
        Открывает месячную таблицу, при необходимости заводя её.

        Новый месяц — новая таблица. Создавать её руками каждый раз
        неудобно и легко забыть, поэтому она делается КОПИЕЙ прошлого
        месяца: так переносятся и вкладки, и оформление, и — главное —
        права доступа. Если создавать с нуля, файл окажется в диске
        сервисного аккаунта и вы его просто не увидите.

        :param want: название нужной таблицы
        :param previous: название прошломесячной, откуда копировать
        """
        try:
            self.sh = self._open_by_title(want)
            log.info("Таблица «%s» найдена", want)
            return self.sh
        except gspread.exceptions.SpreadsheetNotFound:
            pass

        if not previous:
            raise RuntimeError(
                f"Таблица «{want}» не найдена, и копировать не с чего. "
                f"Создайте её вручную и откройте доступ сервисному аккаунту."
            )

        try:
            src = self._open_by_title(previous)
        except gspread.exceptions.SpreadsheetNotFound:
            raise RuntimeError(
                f"Нет ни «{want}», ни «{previous}». "
                f"Создайте таблицу вручную и откройте доступ."
            ) from None

        log.info("Таблицы «%s» нет — копирую с «%s»", want, previous)
        self.sh = self.gc.copy(src.id, title=want, copy_permissions=True)
        log.info("Создана таблица «%s» (доступ унаследован)", want)
        return self.sh

    def archive_to_excel(self, out_path: Path) -> Path:
        """
        Сохраняет все листы таблицы в локальный xlsx.

        Читает значения через Sheets API, а не через выгрузку Drive:
        Drive у вас бывает закрыт файрволом, а Sheets работает всегда.
        Формулы сохраняются как посчитанные значения — для архива это
        то, что нужно: цифры не поедут, если потом изменится источник.
        """
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)

        titles = [w.title for w in self.sh.worksheets()]
        data = []
        for i in range(0, len(titles), 10):
            chunk = titles[i:i + 10]
            data += self.batch_get([self.quote_sheet(t) for t in chunk])

        wb = Workbook()
        wb.remove(wb.active)
        for title, rows in zip(titles, data):
            safe = re.sub(r"[\\/*?:\[\]]", "-", title)[:31]
            ws = wb.create_sheet(safe)
            for r in (rows or []):
                ws.append([("" if c is None else c) for c in r])
            if ws.max_row:
                ws.freeze_panes = "A2"
        wb.save(out_path)
        log.info("Архив сохранён: %s (листов %d)", out_path.name, len(titles))
        return out_path

    def list_date_sheets(self, d_from: date, d_to: date) -> Dict[date, str]:
        """
        Находит вкладки, чьи имена — даты в заданном диапазоне.

        Возвращает {дата: название листа}, отсортировано по дате.
        Вкладки с нечисловыми именами и опечатками в дате отсеиваются
        сами: они либо не разбираются, либо попадают вне диапазона.
        """
        found: Dict[date, str] = {}
        for w in self.sh.worksheets():
            d = DateResolver.parse_sheet_name(w.title)
            if d and d_from <= d <= d_to:
                found[d] = w.title
        return dict(sorted(found.items()))

    def batch_get(self, ranges: List[str]) -> List[List[List[Any]]]:
        """
        Читает несколько диапазонов ОДНИМ запросом к API.

        Так история за месяц обходится в один вызов на таблицу
        вместо тридцати. Пустые диапазоны возвращаются как [].
        """
        if not ranges:
            return []
        res = self.sh.values_batch_get(ranges)
        return [vr.get("values", []) for vr in res.get("valueRanges", [])]

    @staticmethod
    def quote_sheet(title: str) -> str:
        """Оборачивает имя листа в кавычки — иначе точки в дате ломают range."""
        return "'" + str(title).replace("'", "''") + "'"

    def write(self, payload: List[Dict[str, Any]]) -> None:
        """
        Пишет несколько диапазонов одним запросом.

        batch_update экономит квоту API: вместо N вызовов — один.
        value_input_option='USER_ENTERED' нужен, чтобы формулы
        (=CEILING, =SUM) попали в ячейку формулами, а не текстом.
        """
        if not payload:
            return
        self.ws.batch_update(payload, value_input_option="USER_ENTERED")
        log.info("Записано диапазонов: %d", len(payload))

    # ------------------------------------------------------------------
    # ФОРМАТИРОВАНИЕ
    # ------------------------------------------------------------------

    def _sheet_id(self) -> int:
        """Числовой ID листа — нужен для запросов Sheets API."""
        return self.ws._properties["sheetId"]

    def apply_format(self, requests: List[Dict[str, Any]],
                     strict: bool = True) -> None:
        """
        Отправляет пакет запросов форматирования одним вызовом.

        :param strict: False — не падать при ошибке, только предупредить.
                       Нужно для необязательного оформления: данные уже
                       записаны, и терять их из-за косметики не стоит.
        """
        if not requests:
            return
        if DEMO:
            log.info("Демо-режим: оформление в Excel не переносится "
                     "(правил: %d)", len(requests))
            return
        try:
            self.sh.batch_update({"requests": requests})
            log.info("Применено правил форматирования: %d", len(requests))
        except Exception as e:
            if strict:
                raise
            log.warning("Форматирование не применилось (%s): %s",
                        type(e).__name__, str(e)[:200])

    def apply_tables(self, requests: List[Dict[str, Any]]) -> None:
        """
        Создаёт умные таблицы — отдельным вызовом, по одной за раз.

        addTable — новый метод Sheets API, он нестабилен в пакете
        с другими правками того же диапазона и может вернуть 500.
        Поэтому каждая таблица отправляется своим запросом, а сбой
        не роняет скрипт: данные к этому моменту уже на листе.
        """
        for req in requests:
            name = req.get("addTable", {}).get("table", {}).get("name", "?")
            try:
                self.sh.batch_update({"requests": [req]})
                log.info("Создана умная таблица: %s", name)
            except Exception as e:
                log.warning("Умная таблица '%s' не создана (%s). "
                            "Данные и форматирование на месте — "
                            "таблицу можно включить вручную: "
                            "выделить блок -> Формат -> Преобразовать в таблицу",
                            name, type(e).__name__)

    # --- готовые кирпичики форматирования ---

    def req_header(self, row: int, col_start: int, col_end: int,
                   bg: Dict, bold: bool = True,
                   font_color: Optional[Dict] = None) -> Dict:
        """
        Заливка и шрифт строки заголовка.

        Индексы нулевые и полуоткрытые: row=0 — первая строка,
        col_end не включается.
        """
        fmt = {
            "backgroundColor": bg,
            "textFormat": {
                "bold": bold,
                "foregroundColor": font_color or Palette.BLACK,
                "fontSize": 10,
            },
            "horizontalAlignment": "CENTER",
            "verticalAlignment": "MIDDLE",
            "wrapStrategy": "WRAP",
        }
        return {
            "repeatCell": {
                "range": {
                    "sheetId": self._sheet_id(),
                    "startRowIndex": row, "endRowIndex": row + 1,
                    "startColumnIndex": col_start, "endColumnIndex": col_end,
                },
                "cell": {"userEnteredFormat": fmt},
                "fields": "userEnteredFormat(backgroundColor,textFormat,"
                          "horizontalAlignment,verticalAlignment,wrapStrategy)",
            }
        }

    def req_row_fill(self, row: int, col_start: int, col_end: int,
                     bg: Dict, bold: bool = False) -> Dict:
        """Заливка обычной строки данных."""
        return {
            "repeatCell": {
                "range": {
                    "sheetId": self._sheet_id(),
                    "startRowIndex": row, "endRowIndex": row + 1,
                    "startColumnIndex": col_start, "endColumnIndex": col_end,
                },
                "cell": {"userEnteredFormat": {
                    "backgroundColor": bg,
                    "textFormat": {"bold": bold, "fontSize": 10},
                }},
                "fields": "userEnteredFormat(backgroundColor,textFormat)",
            }
        }

    def req_borders(self, row_start: int, row_end: int,
                    col_start: int, col_end: int) -> Dict:
        """Тонкие рамки по всем ячейкам диапазона."""
        line = {"style": "SOLID", "width": 1, "color": Palette.BORDER}
        return {
            "updateBorders": {
                "range": {
                    "sheetId": self._sheet_id(),
                    "startRowIndex": row_start, "endRowIndex": row_end,
                    "startColumnIndex": col_start, "endColumnIndex": col_end,
                },
                "top": line, "bottom": line, "left": line, "right": line,
                "innerHorizontal": line, "innerVertical": line,
            }
        }

    def req_merge(self, row_start: int, row_end: int,
                  col_start: int, col_end: int) -> Dict:
        """Объединение ячеек — для колонки «Часы» по 5 строк проектов."""
        return {
            "mergeCells": {
                "range": {
                    "sheetId": self._sheet_id(),
                    "startRowIndex": row_start, "endRowIndex": row_end,
                    "startColumnIndex": col_start, "endColumnIndex": col_end,
                },
                "mergeType": "MERGE_ALL",
            }
        }

    def req_center(self, row_start: int, row_end: int,
                   col_start: int, col_end: int, bold: bool = False) -> Dict:
        """Выравнивание по центру (по горизонтали и вертикали)."""
        return {
            "repeatCell": {
                "range": {
                    "sheetId": self._sheet_id(),
                    "startRowIndex": row_start, "endRowIndex": row_end,
                    "startColumnIndex": col_start, "endColumnIndex": col_end,
                },
                "cell": {"userEnteredFormat": {
                    "horizontalAlignment": "CENTER",
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {"bold": bold, "fontSize": 10},
                }},
                "fields": "userEnteredFormat(horizontalAlignment,"
                          "verticalAlignment,textFormat)",
            }
        }

    def req_number_format(self, row_start: int, row_end: int,
                          col_start: int, col_end: int,
                          pattern: str = "0.0%",
                          ntype: str = "PERCENT") -> Dict:
        """
        Задаёт числовой формат ячеек — например процентный.

        Для процентов в ячейке должна лежать ДОЛЯ (0.327), тогда
        Sheets покажет 32,7% и цифра останется числом: по ней можно
        строить диаграммы и считать средние.
        """
        return {
            "repeatCell": {
                "range": {
                    "sheetId": self._sheet_id(),
                    "startRowIndex": row_start, "endRowIndex": row_end,
                    "startColumnIndex": col_start, "endColumnIndex": col_end,
                },
                "cell": {"userEnteredFormat": {
                    "numberFormat": {"type": ntype, "pattern": pattern}}},
                "fields": "userEnteredFormat.numberFormat",
            }
        }

    def req_col_width(self, col_start: int, col_end: int, px: int) -> Dict:
        """Ширина колонок в пикселях."""
        return {
            "updateDimensionProperties": {
                "range": {
                    "sheetId": self._sheet_id(),
                    "dimension": "COLUMNS",
                    "startIndex": col_start, "endIndex": col_end,
                },
                "properties": {"pixelSize": px},
                "fields": "pixelSize",
            }
        }

    def req_add_table(self, name: str, row_start: int, row_end: int,
                      col_start: int, col_end: int,
                      column_names: Optional[List[str]] = None) -> Dict:
        """
        Создаёт «Умную таблицу» Google Sheets на диапазоне.

        Умная таблица даёт выпадающие фильтры в шапке, режим фильтрации
        и именованный объект (Таблица39 и т.п.). Первая строка диапазона
        считается заголовком.

        :param name: имя таблицы, видно в углу блока
        :param column_names: имена колонок; нужны, чтобы Google
                             корректно построил фильтры
        """
        spec = {
            "name": name,
            "range": {
                "sheetId": self._sheet_id(),
                "startRowIndex": row_start, "endRowIndex": row_end,
                "startColumnIndex": col_start, "endColumnIndex": col_end,
            },
        }
        if column_names:
            # Только имена: тип колонки Google определяет сам.
            # Явный columnType — частая причина ошибки 500 на addTable.
            spec["columnProperties"] = [
                {"columnIndex": i, "columnName": str(cn)}
                for i, cn in enumerate(column_names)
            ]
        return {"addTable": {"table": spec}}

    # --- условное форматирование ---

    GOOD_GREEN = {"red": 0.05, "green": 0.50, "blue": 0.25}
    BAD_RED = {"red": 0.75, "green": 0.10, "blue": 0.10}

    def req_sign_colors(self, row_start: int, row_end: int,
                        col_start: int, col_end: int,
                        growth_is_bad: bool = True) -> List[Dict]:
        """
        Красит числа по знаку: рост одним цветом, падение другим.

        :param growth_is_bad: True для потерь и просрочек — там рост
            это ухудшение. False для объёма звонков и принятых —
            там рост скорее хорошая новость.

        Возвращает два правила: для положительных и отрицательных.
        """
        up = self.BAD_RED if growth_is_bad else self.GOOD_GREEN
        down = self.GOOD_GREEN if growth_is_bad else self.BAD_RED
        rng = {
            "sheetId": self._sheet_id(),
            "startRowIndex": row_start, "endRowIndex": row_end,
            "startColumnIndex": col_start, "endColumnIndex": col_end,
        }

        def rule(cond_type: str, color: Dict) -> Dict:
            return {
                "addConditionalFormatRule": {
                    "rule": {
                        "ranges": [rng],
                        "booleanRule": {
                            "condition": {
                                "type": cond_type,
                                "values": [{"userEnteredValue": "0"}],
                            },
                            "format": {"textFormat": {
                                "foregroundColor": color, "bold": True}},
                        },
                    },
                    "index": 0,
                }
            }

        return [rule("NUMBER_GREATER", up), rule("NUMBER_LESS", down)]

    def req_color_scale(self, row_start: int, row_end: int,
                        col_start: int, col_end: int,
                        reverse: bool = False) -> Dict:
        """
        Цветовая шкала: от зелёного к красному по величине значения.

        :param reverse: True — наоборот, большое значение зелёное.
        """
        green = {"red": 0.72, "green": 0.88, "blue": 0.75}
        yellow = {"red": 1.0, "green": 0.90, "blue": 0.60}
        red = {"red": 0.96, "green": 0.72, "blue": 0.70}
        lo, hi = (red, green) if reverse else (green, red)
        return {
            "addConditionalFormatRule": {
                "rule": {
                    "ranges": [{
                        "sheetId": self._sheet_id(),
                        "startRowIndex": row_start, "endRowIndex": row_end,
                        "startColumnIndex": col_start, "endColumnIndex": col_end,
                    }],
                    "gradientRule": {
                        "minpoint": {"color": lo, "type": "MIN"},
                        "midpoint": {"color": yellow, "type": "PERCENTILE",
                                     "value": "50"},
                        "maxpoint": {"color": hi, "type": "MAX"},
                    },
                },
                "index": 0,
            }
        }

    def req_freeze(self, rows: int = 1) -> Dict:
        """Закрепление верхних строк, чтобы шапка не уезжала при прокрутке."""
        return {
            "updateSheetProperties": {
                "properties": {
                    "sheetId": self._sheet_id(),
                    "gridProperties": {"frozenRowCount": rows},
                },
                "fields": "gridProperties.frozenRowCount",
            }
        }


# ==========================================================================
# КЭШ ИСТОРИИ
# ==========================================================================

class HistoryCache:
    """
    Локальный кэш разобранной истории — чтобы не перечитывать
    прошлые листы при каждом запуске.

    Хранится обычным CSV в папке data/cache: файл открывается
    Excel'ом, при сомнениях его можно просто удалить — тогда
    следующий запуск соберёт всё заново.

    Ключ записи — пара (Дата, Источник). Уже закэшированные дни
    из Google Sheets повторно не читаются.
    """

    def __init__(self, path: Path, key_cols: Sequence[str] = ("Дата", "Источник")):
        self.path = Path(path)
        self.key_cols = list(key_cols)

    def load(self) -> pd.DataFrame:
        """Читает кэш. Отсутствующий или битый файл -> пустая таблица."""
        if not self.path.exists():
            return pd.DataFrame()
        try:
            df = pd.read_csv(self.path)
            if "Дата" in df.columns:
                df["Дата"] = pd.to_datetime(df["Дата"], errors="coerce").dt.date
            log.info("Кэш %s: %d строк", self.path.name, len(df))
            return df
        except Exception as e:
            log.warning("Кэш %s не прочитан (%s) — соберу заново",
                        self.path.name, type(e).__name__)
            return pd.DataFrame()

    def cached_dates(self, source: str) -> set:
        """Какие дни по этому источнику уже лежат в кэше."""
        df = self.load()
        if df.empty or "Источник" not in df.columns:
            return set()
        return set(df.loc[df["Источник"] == source, "Дата"].dropna())

    def merge(self, fresh: pd.DataFrame) -> pd.DataFrame:
        """
        Дописывает свежие дни к кэшу и сохраняет.

        При совпадении ключа побеждает свежая запись — так
        перезалитый задним числом день попадёт в кэш обновлённым.
        """
        old = self.load()
        if fresh.empty:
            return old
        if old.empty:
            merged = fresh.copy()
        else:
            merged = pd.concat([old, fresh], ignore_index=True)
            subset = [c for c in self.key_cols if c in merged.columns]
            extra = [c for c in ("Час", "Проект", "Оператор") if c in merged.columns]
            merged = merged.drop_duplicates(subset=subset + extra, keep="last")

        self.path.parent.mkdir(parents=True, exist_ok=True)
        merged.to_csv(self.path, index=False)
        log.info("Кэш %s обновлён: %d строк", self.path.name, len(merged))
        return merged


# ==========================================================================
# ВАЛИДАЦИЯ
# ==========================================================================

class Validator:
    """Сверка агрегатов с первоисточником. При расхождении — остановка."""

    @staticmethod
    def check_total(aggregated: float, source: float, label: str,
                    strict: bool = True) -> None:
        """
        Сравнивает сумму после группировки с исходным итогом.

        atol=1e-9 гасит бинарную погрешность float, но реальное
        расхождение хотя бы в единицу выявит.
        """
        ok = np.isclose(aggregated, source, rtol=0, atol=1e-9)
        if ok:
            log.info("Сверка '%s' пройдена: %s", label, aggregated)
            return

        msg = (f"РАСХОЖДЕНИЕ '{label}': после группировки={aggregated}, "
               f"в источнике={source}, дельта={aggregated - source}")
        if strict:
            raise AssertionError(msg)
        log.warning(msg)

    @staticmethod
    def report(pairs: Sequence[tuple]) -> pd.DataFrame:
        """Собирает таблицу сверок для визуального контроля перед выгрузкой."""
        rows = [{"Показатель": lbl, "Расчёт": agg, "Источник": src,
                 "Дельта": agg - src, "OK": "✓" if agg == src else "✗"}
                for lbl, agg, src in pairs]
        return pd.DataFrame(rows)


def preview(df: pd.DataFrame, title: str) -> None:
    """
    Печатает датафрейм в консоль DataSpell перед выгрузкой.

    Нужно для визуального контроля: сначала смотрите цифры,
    потом запускаете ячейку с выгрузкой в облако.
    """
    print("\n" + "=" * 78)
    print(title)
    print("=" * 78)
    try:
        from IPython.display import display as _d
        _d(df)
    except ImportError:
        print(df.to_string())


# ==========================================================================
# ДЕМО-РЕЖИМ
# ==========================================================================
#
# DEMO=1 в .env — таблицы Google подменяются Excel-книгами в
# data/demo_output. Так ноутбуки запускаются без доступа к Google,
# на демо-выгрузках из generate_etl_demo.py.
#
# Как устроено. SheetsLoader и ноутбуки работают с таблицей через
# небольшой набор методов gspread: worksheets(), add_worksheet(),
# values_batch_get() и ещё несколько. DemoBook и DemoSheet повторяют
# ровно эти методы, только поверх файла Excel. Поэтому код ноутбуков
# в демо-режиме тот же: connect() открывает книгу вместо таблицы,
# дальше всё идёт как обычно.
#
# Чего в демо нет:
#   - оформления: заливки, рамки, объединения — это запросы Sheets API,
#     в Excel они пропускаются;
#   - вычисления формул: книгу пишет openpyxl, а он формулы не считает.
#     Поэтому формула сохраняется ТЕКСТОМ — видно, что ушло бы в Google, —
#     и при чтении обратно приходит текстом, а не числом.


class DemoSheet:
    """Лист демо-книги: те методы листа gspread, что вызывает код."""

    def __init__(self, book: "DemoBook", xl_sheet):
        self.book = book
        self.xl = xl_sheet            # лист openpyxl
        # Номер листа нужен запросам оформления: в демо они не
        # выполняются, но собираются так же, как для Google
        self._properties = {"sheetId": book.xl.worksheets.index(xl_sheet)}

    @property
    def title(self) -> str:
        return self.xl.title

    def get(self, rng: str) -> List[List[str]]:
        return self.book.read(self.xl, rng)

    def batch_update(self, data: List[Dict[str, Any]],
                     value_input_option: str = "") -> None:
        for item in data:
            self.book.put(self.xl, item["range"], item["values"])
        self.book.save()

    def batch_clear(self, ranges: List[str]) -> None:
        for rng in ranges:
            self.book.clear(self.xl, rng)
        self.book.save()

    def clear(self) -> None:
        self.xl.delete_rows(1, self.xl.max_row)
        self.book.save()

    def format(self, *args, **kwargs) -> None:
        """Числовой формат ячейки — оформление, в демо пропускается."""


class DemoBook:
    """Excel-книга в data/demo_output вместо таблицы Google."""

    def __init__(self, title: str):
        self.title = title
        self.path = DEMO_DIR / f"{title}.xlsx"
        if self.path.exists():
            self.xl = openpyxl.load_workbook(self.path)
        else:
            # Новая таблица Google создаётся с одним пустым листом —
            # здесь так же. Файл появится при первой записи.
            self.xl = openpyxl.Workbook()
            self.xl.active.title = "Лист1"

    def save(self) -> None:
        DEMO_DIR.mkdir(parents=True, exist_ok=True)
        self.xl.save(self.path)

    # --- листы ---

    def worksheets(self) -> List[DemoSheet]:
        return [DemoSheet(self, s) for s in self.xl.worksheets]

    def worksheet(self, name: str) -> DemoSheet:
        if name not in self.xl.sheetnames:
            raise gspread.exceptions.WorksheetNotFound(name)
        return DemoSheet(self, self.xl[name])

    def add_worksheet(self, title: str, rows: int = 0,
                      cols: int = 0) -> DemoSheet:
        sheet = DemoSheet(self, self.xl.create_sheet(title))
        self.save()
        return sheet

    def del_worksheet(self, sheet: DemoSheet) -> None:
        self.xl.remove(sheet.xl)
        self.save()

    # --- значения ---

    def _locate(self, full_range: str) -> tuple:
        """«'Итог'!C5» -> (лист openpyxl, «C5»)."""
        name, rng = full_range.rsplit("!", 1)
        name = name.strip("'").replace("''", "'")
        if name not in self.xl.sheetnames:
            # Google на такой диапазон тоже отвечает ошибкой
            raise ValueError(f"Демо-режим: в книге {self.path.name} "
                             f"нет листа «{name}»")
        return self.xl[name], rng

    def values_batch_get(self, ranges: List[str]) -> Dict[str, Any]:
        found = [self.read(*self._locate(r)) for r in ranges]
        return {"valueRanges": [{"values": v} for v in found]}

    def values_batch_update(self, body: Dict[str, Any]) -> None:
        for item in body["data"]:
            sheet, rng = self._locate(item["range"])
            self.put(sheet, rng, item["values"])
        self.save()

    def batch_update(self, body: Dict[str, Any]) -> None:
        """Запросы Sheets API — оформление и порядок вкладок. Пропускаем."""

    def fetch_sheet_metadata(self) -> Dict[str, Any]:
        """Описание листов: reset_sheet_format ищет в нём правила оформления."""
        return {"sheets": [{"properties": {"sheetId": i}}
                           for i in range(len(self.xl.worksheets))]}

    # --- ячейки ---

    @staticmethod
    def put(sheet, rng: str, values: List[List[Any]]) -> None:
        """Пишет блок значений, начиная с левой верхней ячейки диапазона."""
        col0, row0 = range_boundaries(rng)[:2]
        for i, row in enumerate(values):
            for j, v in enumerate(row):
                cell = sheet.cell(row=row0 + i, column=col0 + j)
                cell.value = None if v == "" else v
                # Формулу Google храним текстом: посчитать её openpyxl
                # не может, а записанная формулой она испортит файл —
                # CEILING с одним аргументом и СУММЕСЛИ Excel не знает
                if isinstance(v, str) and v.startswith("="):
                    cell.data_type = "s"

    @staticmethod
    def _bounds(sheet, rng: str) -> tuple:
        """
        Границы диапазона, обрезанные по заполненной части листа.

        openpyxl заводит ячейку на каждое обращение, и чтение A1:GZ1
        без обрезки раздуло бы файл пустыми ячейками.
        """
        c1, r1, c2, r2 = range_boundaries(rng)
        return c1, r1, min(c2, sheet.max_column), min(r2, sheet.max_row)

    def read(self, sheet, rng: str) -> List[List[str]]:
        """
        Значения диапазона так, как их отдаёт Sheets API: строками,
        без пустых ячеек в конце строки и без пустых строк в конце.
        """
        c1, r1, c2, r2 = self._bounds(sheet, rng)
        rows = []
        for r in range(r1, r2 + 1):
            row = []
            for c in range(c1, c2 + 1):
                v = sheet.cell(row=r, column=c).value
                row.append("" if v is None else str(v))
            while row and row[-1] == "":
                row.pop()
            rows.append(row)
        while rows and not rows[-1]:
            rows.pop()
        return rows

    def clear(self, sheet, rng: str) -> None:
        c1, r1, c2, r2 = self._bounds(sheet, rng)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                sheet.cell(row=r, column=c).value = None
